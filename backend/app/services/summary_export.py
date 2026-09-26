"""图谱总结导出：Markdown 文档 / Word 文档 / PPTX 演示

PPT 里有一页是**用形状与箭头原生画出来的流程图**（不是贴图），
保证在任何机器上打开都不糊、可继续编辑。
"""
import os
from datetime import datetime
from typing import Dict, List, Optional

from loguru import logger

from app.services.graph_summary import LEVEL_LABELS, RELATION_LABELS, to_ai_digest, to_mermaid

# 与前端主题一致的配色（雾霾蓝 / 薄荷 / 暖杏 / 暖白）
C_PRIMARY = (0x4F, 0x6F, 0x8F)
C_PRIMARY_DARK = (0x3D, 0x5A, 0x75)
C_MINT = (0x7C, 0xB8, 0xA0)
C_APRICOT = (0xD4, 0xA5, 0x74)
C_LILAC = (0xB9, 0xA7, 0xCC)
C_TEXT = (0x2C, 0x3E, 0x4F)
C_MUTED = (0x5E, 0x72, 0x86)
C_BG = (0xF5, 0xF0, 0xEB)
C_WHITE = (0xFF, 0xFF, 0xFF)

LEVEL_RGB = {1: C_PRIMARY, 2: C_MINT, 3: C_APRICOT, 4: C_LILAC}

EXPORT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__)))), "exports")


def _ensure_dir() -> str:
    os.makedirs(EXPORT_DIR, exist_ok=True)
    # 只保留最近 30 个导出文件，避免长期运行堆积
    try:
        files = [os.path.join(EXPORT_DIR, f) for f in os.listdir(EXPORT_DIR)]
        files = [f for f in files if os.path.isfile(f)]
        files.sort(key=os.path.getmtime, reverse=True)
        for old in files[30:]:
            os.remove(old)
    except OSError as exc:
        # 清理旧产物失败（如文件被打开占用）不应影响导出本身
        logger.debug(f"清理旧导出产物失败：{type(exc).__name__}: {exc}")
    return EXPORT_DIR


def export_path(stem: str, ext: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return os.path.join(_ensure_dir(), f"{stem}_{stamp}.{ext}")


# ---------------------------------------------------------------- Markdown


def to_markdown(summary: Dict, mermaid: Optional[str] = None, title: str = "知识图谱总结") -> str:
    if summary.get("empty"):
        return f"# {title}\n\n图谱暂无知识点。\n"

    ov = summary["overview"]
    md: List[str] = [
        f"# {title}",
        "",
        f"> 生成时间：{summary.get('generated_at', '')}　|　"
        f"知识点 {ov['node_count']}　关联 {ov['link_count']}　文件 {ov['file_count']}",
        "",
        "## 一、总览",
        "",
        "| 指标 | 数值 | 指标 | 数值 |",
        "|---|---|---|---|",
        f"| 知识点 | {ov['node_count']} | 关联连线 | {ov['link_count']} |",
        f"| 来源文件 | {ov['file_count']} | 领域数 | {ov['domain_count']} |",
        f"| 知识簇 | {ov['cluster_count']} | 孤立知识点 | {ov['isolated_count']} |",
        f"| 平均度 | {ov['avg_degree']} | 图谱密度 | {ov['density']} |",
        f"| 知识库锚定概念 | {ov['kb_anchor_count']} | 知识库桥接 | {ov['kb_bridge_count']}（跨文件 {ov['cross_file_bridge_count']}） |",
        "",
        "**关系类型分布**：" + "、".join(f"{k} {v}" for k, v in ov["relation_distribution"].items()),
        "",
        "## 二、知识层级结构",
        "",
        "| 层级 | 名称 | 数量 | 代表知识点 |",
        "|---|---|---|---|",
    ]
    for lv in summary["levels"]:
        md.append(f"| L{lv['level']} | {lv['label']} | {lv['count']} | {'、'.join(lv['samples'][:6])} |")

    md += ["", "## 三、领域分布", "", "| 领域 | 知识点数 | 占比 | 平均层级 | 代表知识点 |", "|---|---|---|---|---|"]
    for d in summary["domains"]:
        md.append(f"| {d['domain']} | {d['count']} | {d['share'] * 100:.1f}% | L{d['avg_level']} | "
                  f"{'、'.join(d['top_nodes'][:4])} |")

    if summary["paths"]:
        md += ["", "## 四、学习主线（把图谱拉成流程）", ""]
        for i, p in enumerate(summary["paths"], 1):
            md.append(f"### 主线 {i}：{p['title']}（{p['length']} 步）")
            md.append("")
            for s in p["steps"]:
                rel = f"　←（{s['relation_label_from_prev']}）" if s["relation_label_from_prev"] else ""
                md.append(f"{s['order']}. **{s['title']}**　`L{s['level']} {s['level_label']}`　"
                          f"`{s['domain']}`{rel}")
            md.append("")

    if summary["hubs"]:
        md += ["## 五、核心枢纽知识点", "", "| 知识点 | 层级 | 领域 | 连接数 | 连到谁 |", "|---|---|---|---|---|"]
        for h in summary["hubs"][:10]:
            md.append(f"| {h['title']} | L{h['level']} | {h['domain']} | {h['degree']} | "
                      f"{'、'.join(h['connected_titles'][:5])} |")
        md.append("")

    if summary["clusters"]:
        md += ["## 六、知识簇", ""]
        for i, c in enumerate(summary["clusters"], 1):
            md.append(f"- **簇 {i}｜{c['label']}**　{c['size']} 个知识点 · "
                      f"主领域 {c['dominant_domain']} · 涉及 {c['file_count']} 个文件")
            md.append(f"  - 代表节点：{'、'.join(n['title'] for n in c['top_nodes'])}")
        md.append("")

    if summary["bridges"]:
        md += ["## 七、知识库桥接（由知识库推理建立的知识通路）", "",
               "| 源 | 关系 | 目标 | 分数 | 证据 |", "|---|---|---|---|---|"]
        for b in summary["bridges"][:15]:
            md.append(f"| {b['source']} | {b['relation_label']} | {b['target']} | {b['score']} | "
                      f"{b['evidence'] or b['source_text'] or '—'} |")
        md.append("")

    if summary["isolated"]:
        md += ["## 八、孤立知识点（待补全）", ""]
        md.extend(f"- {n['title']}（L{n['level']} {n['level_label']}，{n['domain']}）"
                  for n in summary["isolated"][:20])
        md.append("")

    md += ["## 九、结论与建议", ""]
    md.extend(f"{i}. {s}" for i, s in enumerate(summary["suggestions"], 1))

    if mermaid:
        md += ["", "## 十、总结流程图（Mermaid，可交给 AI 继续学习）", "",
               "```mermaid", mermaid.strip(), "```", ""]
    return "\n".join(md)


# ---------------------------------------------------------------- Word


def _set_cjk(run, name: str = "微软雅黑"):
    from docx.oxml.ns import qn
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def to_docx(summary: Dict, mermaid: Optional[str] = None, title: str = "知识图谱总结") -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    path = export_path("知识图谱总结", "docx")
    doc = Document()

    # 全局样式：中文用微软雅黑
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(
        __import__("docx.oxml.ns", fromlist=["qn"]).qn("w:eastAsia"), "微软雅黑")

    h = doc.add_heading(title, level=0)
    for r in h.runs:
        _set_cjk(r)
        r.font.color.rgb = RGBColor(*C_PRIMARY_DARK)

    sub = doc.add_paragraph()
    _set_cjk(sub.add_run(
        f"生成时间：{summary.get('generated_at', '')}　|　"
        f"知识点 {summary['overview']['node_count']}　关联 {summary['overview']['link_count']}　"
        f"文件 {summary['overview']['file_count']}"), )
    sub.runs[0].font.size = Pt(9)
    sub.runs[0].font.color.rgb = RGBColor(*C_MUTED)
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if summary.get("empty"):
        doc.add_paragraph("图谱暂无知识点。")
        doc.save(path)
        return path

    ov = summary["overview"]

    def heading(text, level=1):
        hh = doc.add_heading(text, level=level)
        for r in hh.runs:
            _set_cjk(r)
            r.font.color.rgb = RGBColor(*C_PRIMARY_DARK)
        return hh

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, htxt in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(htxt)
            _set_cjk(run)
            run.bold = True
            run.font.size = Pt(9.5)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(val))
                _set_cjk(run)
                run.font.size = Pt(9)
        return t

    heading("一、总览")
    table(["指标", "数值", "指标", "数值"], [
        ["知识点", ov["node_count"], "关联连线", ov["link_count"]],
        ["来源文件", ov["file_count"], "领域数", ov["domain_count"]],
        ["知识簇", ov["cluster_count"], "孤立知识点", ov["isolated_count"]],
        ["平均度", ov["avg_degree"], "图谱密度", ov["density"]],
        ["知识库锚定概念", ov["kb_anchor_count"],
         "知识库桥接", f"{ov['kb_bridge_count']}（跨文件 {ov['cross_file_bridge_count']}）"],
    ])
    p = doc.add_paragraph()
    _set_cjk(p.add_run("关系类型分布：" + "、".join(
        f"{k} {v}" for k, v in ov["relation_distribution"].items())))

    heading("二、知识层级结构")
    table(["层级", "名称", "数量", "代表知识点"],
          [[f"L{lv['level']}", lv["label"], lv["count"], "、".join(lv["samples"][:6])]
           for lv in summary["levels"]])

    heading("三、领域分布")
    table(["领域", "知识点数", "占比", "平均层级", "代表知识点"],
          [[d["domain"], d["count"], f"{d['share'] * 100:.1f}%", f"L{d['avg_level']}",
            "、".join(d["top_nodes"][:4])] for d in summary["domains"]])

    if summary["paths"]:
        heading("四、学习主线（把图谱拉成流程）")
        for i, p_ in enumerate(summary["paths"], 1):
            hp = doc.add_paragraph()
            _set_cjk(hp.add_run(f"主线 {i}：{p_['title']}"), )
            hp.runs[0].bold = True
            for s in p_["steps"]:
                rel = f"　←（{s['relation_label_from_prev']}）" if s["relation_label_from_prev"] else ""
                doc.add_paragraph(
                    f"{s['order']}. {s['title']}　[L{s['level']} {s['level_label']}]　"
                    f"{s['domain']}{rel}", style="List Number")

    if summary["hubs"]:
        heading("五、核心枢纽知识点")
        table(["知识点", "层级", "领域", "连接数", "连到谁"],
              [[h["title"], f"L{h['level']}", h["domain"], h["degree"],
                "、".join(h["connected_titles"][:5])] for h in summary["hubs"][:10]])

    if summary["clusters"]:
        heading("六、知识簇")
        for i, c in enumerate(summary["clusters"], 1):
            doc.add_paragraph(f"簇 {i}｜{c['label']}　{c['size']} 个知识点 · "
                              f"主领域 {c['dominant_domain']} · {c['file_count']} 个文件",
                              style="List Bullet")
            doc.add_paragraph("代表节点：" + "、".join(n["title"] for n in c["top_nodes"]),
                              style="List Bullet 2")

    if summary["bridges"]:
        heading("七、知识库桥接")
        table(["源", "关系", "目标", "分数", "证据"],
              [[b["source"], b["relation_label"], b["target"], b["score"],
                b["evidence"] or b["source_text"] or "—"] for b in summary["bridges"][:15]])

    if summary["isolated"]:
        heading("八、孤立知识点（待补全）")
        for n in summary["isolated"][:20]:
            doc.add_paragraph(
                f"{n['title']}（L{n['level']} {n['level_label']}，{n['domain']}）",
                style="List Bullet")

    heading("九、结论与建议")
    for s in summary["suggestions"]:
        doc.add_paragraph(s, style="List Number")

    if mermaid:
        heading("十、总结流程图（Mermaid，可交给 AI 继续学习）")
        mp = doc.add_paragraph()
        run = mp.add_run(mermaid.strip())
        run.font.name = "Consolas"
        run.font.size = Pt(7.5)

    doc.save(path)
    logger.info(f"总结文档已导出：{path}")
    return path


# ---------------------------------------------------------------- PPTX


def to_pptx(summary: Dict, mermaid: Optional[str] = None,
            title: str = "知识图谱总结") -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Emu, Inches, Pt

    path = export_path("知识图谱总结", "pptx")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def rgb(c):
        return RGBColor(*c)

    def add_rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.12):
        shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                shp.adjustments[0] = radius
            except (IndexError, KeyError) as exc:
                # 某些形状不支持圆角调整值，退化为直角即可（不影响内容）
                logger.debug(f"圆角调整值设置失败（{type(exc).__name__}），按直角绘制")
        return shp

    def add_text(slide, x, y, w, h, text, size=14, color=C_TEXT, bold=False,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.25):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        lines = text.split("\n")
        for i, ln in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.line_spacing = line_spacing
            run = para.add_run()
            run.text = ln
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
            run.font.name = "微软雅黑"
        return box

    def slide_with_title(title_text, kicker=None):
        s = prs.slides.add_slide(blank)
        add_rect(s, 0, 0, 13.333, 0.12, C_PRIMARY, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, 0.6, 0.34, 12.2, 0.6, title_text, size=24, bold=True,
                 color=C_PRIMARY_DARK)
        if kicker:
            add_text(s, 0.62, 0.98, 12.2, 0.34, kicker, size=11, color=C_MUTED)
        return s

    ov = summary.get("overview", {})

    # ---- 1 封面 ----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, 13.333, 7.5, C_BG, shape=MSO_SHAPE.RECTANGLE)
    add_rect(s, 0, 0, 13.333, 0.5, C_PRIMARY, shape=MSO_SHAPE.RECTANGLE)
    add_rect(s, 0, 7.0, 13.333, 0.5, C_MINT, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, 1.0, 2.3, 11.3, 1.2, title, size=44, bold=True, color=C_PRIMARY_DARK)
    add_text(s, 1.02, 3.6, 11.3, 0.5,
             f"知识点 {ov.get('node_count', 0)} · 关联 {ov.get('link_count', 0)} · "
             f"文件 {ov.get('file_count', 0)} · 知识簇 {ov.get('cluster_count', 0)}",
             size=16, color=C_MUTED)
    add_text(s, 1.02, 4.15, 11.3, 0.4,
             f"生成时间 {summary.get('generated_at', '')}　|　"
             f"知识库锚定 {ov.get('kb_anchor_count', 0)} 个概念、"
             f"桥接 {ov.get('kb_bridge_count', 0)} 条知识通路",
             size=12, color=C_MUTED)
    add_text(s, 1.02, 4.9, 11.3, 0.4, "本页内容由知识图谱自动总结，随图谱实时生成", size=11,
             color=C_APRICOT)

    if summary.get("empty"):
        prs.save(path)
        return path

    # ---- 2 总览 ----
    s = slide_with_title("总览", "图谱规模与结构健康度")
    cards = [
        ("知识点", ov.get("node_count", 0), "个"),
        ("关联连线", ov.get("link_count", 0), "条"),
        ("来源文件", ov.get("file_count", 0), "个"),
        ("领域", ov.get("domain_count", 0), "个"),
        ("知识簇", ov.get("cluster_count", 0), "个"),
        ("孤立点", ov.get("isolated_count", 0), "个"),
    ]
    for i, (label, val, unit) in enumerate(cards):
        x = 0.6 + i * 2.06
        add_rect(s, x, 1.6, 1.9, 1.5, C_WHITE, line=C_BG)
        add_text(s, x + 0.1, 1.78, 1.7, 0.6, str(val), size=30, bold=True,
                 color=C_PRIMARY, align=PP_ALIGN.CENTER)
        add_text(s, x + 0.1, 2.48, 1.7, 0.3, label + f"（{unit}）", size=11,
                 color=C_MUTED, align=PP_ALIGN.CENTER)
    add_text(s, 0.62, 3.35, 12.1, 0.35, "关系类型分布", size=14, bold=True, color=C_TEXT)
    rels = list(ov.get("relation_distribution", {}).items())[:8]
    rel_text = "　".join(f"{k} {v}" for k, v in rels) or "—"
    add_text(s, 0.62, 3.75, 12.1, 0.6, rel_text, size=12, color=C_MUTED)
    add_text(s, 0.62, 4.5, 6.0, 2.2,
             "结构指标\n"
             f"平均度 {ov.get('avg_degree', 0)}　图谱密度 {ov.get('density', 0)}\n"
             f"知识库锚定概念 {ov.get('kb_anchor_count', 0)} 个\n"
             f"知识库桥接 {ov.get('kb_bridge_count', 0)} 条"
             f"（跨文件 {ov.get('cross_file_bridge_count', 0)} 条）",
             size=12, color=C_TEXT)
    add_text(s, 6.9, 4.5, 5.8, 2.2,
             "怎么读这份总结\n"
             "① 先看学习主线 → 得到阅读顺序\n"
             "② 再看核心枢纽 → 知道先掌握谁\n"
             "③ 最后看桥接与孤立点 → 知道缺口在哪",
             size=12, color=C_MUTED)

    # ---- 3 层级结构 ----
    s = slide_with_title("知识层级结构", "L1 元概念 → L2 核心理论 → L3 应用实践 → L4 实现工具")
    for i, lv in enumerate(summary["levels"]):
        x = 0.6 + i * 3.14
        add_rect(s, x, 1.6, 2.94, 0.72, LEVEL_RGB.get(lv["level"], C_PRIMARY))
        add_text(s, x + 0.16, 1.72, 2.6, 0.5,
                 f"L{lv['level']} {lv['label']}　{lv['count']} 个", size=15, bold=True,
                 color=C_WHITE)
        add_rect(s, x, 2.32, 2.94, 3.9, C_WHITE, line=C_BG)
        add_text(s, x + 0.16, 2.46, 2.62, 3.6,
                 "\n".join(f"· {s_}" for s_ in lv["samples"][:9]) or "—",
                 size=11, color=C_TEXT, line_spacing=1.4)

    # ---- 4 领域分布 ----
    s = slide_with_title("领域分布", "各领域的知识点规模与代表内容")
    max_share = max([d["share"] for d in summary["domains"]] + [0.01])
    for i, d in enumerate(summary["domains"][:10]):
        y = 1.65 + i * 0.53
        add_text(s, 0.62, y, 2.0, 0.34, d["domain"], size=11, bold=True, color=C_TEXT)
        bar_w = 6.6 * (d["share"] / max_share)
        add_rect(s, 2.7, y + 0.06, max(0.12, bar_w), 0.24,
                 C_PRIMARY if i % 2 == 0 else C_MINT, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, 9.45, y, 1.0, 0.34, f"{d['count']} 个", size=11, color=C_MUTED)
        add_text(s, 10.5, y, 2.4, 0.34, f"L{d['avg_level']}｜" + "、".join(d["top_nodes"][:2]),
                 size=9.5, color=C_MUTED)

    # ---- 5 核心枢纽 ----
    s = slide_with_title("核心枢纽知识点", "连接数最高的知识点 —— 先掌握它们收益最大")
    rows = summary["hubs"][:7]
    if rows:
        tbl = s.shapes.add_table(len(rows) + 1, 5, Inches(0.6), Inches(1.6),
                                 Inches(12.1), Inches(0.4 * (len(rows) + 1))).table
        headers = ["知识点", "层级", "领域", "连接数", "连到谁"]
        widths = [3.2, 1.1, 1.8, 1.2, 4.8]
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        for j, htxt in enumerate(headers):
            cell = tbl.cell(0, j)
            cell.text = htxt
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
                    run.font.bold = True
                    run.font.color.rgb = rgb(C_WHITE)
                    run.font.name = "微软雅黑"
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(C_PRIMARY)
        for i, h in enumerate(rows, 1):
            values = [h["title"], f"L{h['level']}", h["domain"], h["degree"],
                      "、".join(h["connected_titles"][:5])]
            for j, val in enumerate(values):
                cell = tbl.cell(i, j)
                cell.text = str(val)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.font.color.rgb = rgb(C_TEXT)
                        run.font.name = "微软雅黑"

    # ---- 6 知识簇 ----
    s = slide_with_title("知识簇", "彼此连通的知识群 —— 每个簇就是一块可独立学习的模块")
    for i, c in enumerate(summary["clusters"][:6]):
        x = 0.6 + (i % 2) * 6.2
        y = 1.65 + (i // 2) * 1.85
        add_rect(s, x, y, 5.95, 1.65, C_WHITE, line=C_BG)
        add_rect(s, x, y, 0.09, 1.65, C_MINT, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, x + 0.25, y + 0.12, 5.5, 0.34,
                 f"簇 {i + 1}｜{c['label'][:22]}", size=13, bold=True, color=C_PRIMARY_DARK)
        add_text(s, x + 0.25, y + 0.5, 5.5, 0.34,
                 f"{c['size']} 个知识点 · 主领域 {c['dominant_domain']} · "
                 f"{c['file_count']} 个文件", size=10, color=C_MUTED)
        add_text(s, x + 0.25, y + 0.85, 5.5, 0.7,
                 "代表：" + "、".join(n["title"] for n in c["top_nodes"][:5]),
                 size=10, color=C_TEXT)

    # ---- 7 学习主线（原生流程图页） ----
    s = slide_with_title("总结流程图（学习主线）", "把网状知识拉成一条可执行的流程：按箭头方向学")
    y = 1.7
    for pi, p in enumerate(summary["paths"][:3]):
        steps = p["steps"][:6]
        add_text(s, 0.6, y - 0.02, 12.1, 0.3,
                 f"主线 {pi + 1}：{p['length']} 步　{' · '.join(p['domains'][:4])}",
                 size=11, bold=True, color=C_MUTED)
        box_w, gap = 1.72, 0.24
        for si, st in enumerate(steps):
            x = 0.6 + si * (box_w + gap)
            shp = add_rect(s, x, y + 0.34, box_w, 1.0, LEVEL_RGB.get(st["level"], C_PRIMARY),
                           radius=0.14)
            tf = shp.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            run.text = f"{st['order']}. {st['title'][:16]}\nL{st['level']} · {st['domain']}"
            run.font.size = Pt(9.5)
            run.font.color.rgb = rgb(C_WHITE)
            run.font.name = "微软雅黑"
            if si < len(steps) - 1:
                arrow = add_rect(s, x + box_w + 0.02, y + 0.7, gap - 0.04, 0.28,
                                 C_APRICOT, shape=MSO_SHAPE.RIGHT_ARROW)
                if st["relation_label_from_prev"]:
                    add_text(s, x + box_w - 0.05, y + 1.36, gap + 0.5, 0.24,
                             steps[si + 1]["relation_label_from_prev"] or "", size=8,
                             color=C_MUTED, align=PP_ALIGN.CENTER)
        y += 1.72
    add_text(s, 0.6, 6.9, 12.1, 0.4,
             "箭头上的文字＝两个知识点之间的真实关系类型（前置/理论/实现/应用…），"
             "全部来自图谱里的真实连线，可直接作为学习计划或交给 AI 继续展开。",
             size=10, color=C_MUTED)

    # ---- 8 知识库桥接 ----
    if summary["bridges"]:
        s = slide_with_title("知识库桥接", "由知识库推理建立的知识通路（虚线关系）—— 跨文件、跨领域")
        rows = summary["bridges"][:7]
        tbl = s.shapes.add_table(len(rows) + 1, 5, Inches(0.6), Inches(1.6),
                                 Inches(12.1), Inches(0.4 * (len(rows) + 1))).table
        headers = ["源知识点", "关系", "目标知识点", "分数", "证据"]
        for j, w in enumerate([2.7, 1.3, 2.7, 0.9, 4.5]):
            tbl.columns[j].width = Inches(w)
        for j, htxt in enumerate(headers):
            cell = tbl.cell(0, j)
            cell.text = htxt
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
                    run.font.bold = True
                    run.font.color.rgb = rgb(C_WHITE)
                    run.font.name = "微软雅黑"
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(C_MINT)
        for i, b in enumerate(rows, 1):
            values = [b["source"], b["relation_label"], b["target"], b["score"],
                      (b["evidence"] or b["source_text"] or "—")[:60]]
            for j, val in enumerate(values):
                cell = tbl.cell(i, j)
                cell.text = str(val)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.font.color.rgb = rgb(C_TEXT)
                        run.font.name = "微软雅黑"

    # ---- 9 待补全与建议 ----
    s = slide_with_title("待补全与建议", "图谱的短板在哪，下一步补什么")
    left = summary["isolated"][:8]
    add_text(s, 0.62, 1.6, 5.8, 0.35, f"孤立知识点（{len(summary['isolated'])} 个）",
             size=14, bold=True, color=C_APRICOT)
    add_text(s, 0.62, 2.0, 5.8, 4.4,
             "\n".join(f"· {n['title']}（L{n['level']} {n['domain']}）" for n in left) or "无",
             size=11, color=C_TEXT, line_spacing=1.4)
    add_text(s, 6.9, 1.6, 5.8, 0.35, "结论与建议", size=14, bold=True, color=C_PRIMARY)
    add_text(s, 6.9, 2.0, 5.8, 4.4,
             "\n".join(f"{i}. {s_}" for i, s_ in enumerate(summary["suggestions"], 1)),
             size=11.5, color=C_TEXT, line_spacing=1.4)

    # ---- 10 交给 AI ----
    s = slide_with_title("把这张图交给 AI", "总结流程图是机器可读的，直接粘贴给大模型即可继续学习")
    add_text(s, 0.62, 1.6, 12.1, 0.4,
             "在「图谱总结」页面点击「复制 AI 摘要 / 复制 Mermaid」，粘贴到任意支持 Mermaid 的 "
             "对话框或编辑器（ChatGPT / Claude / 飞书 / Notion / VSCode）即可。",
             size=12, color=C_TEXT)
    add_rect(s, 0.6, 2.3, 12.1, 4.3, C_WHITE, line=C_BG)
    add_text(s, 0.85, 2.5, 11.6, 3.9,
             (mermaid or "（未生成流程图）").strip()[:1400],
             size=8, color=C_MUTED, line_spacing=1.15)

    prs.save(path)
    logger.info(f"总结 PPT 已导出：{path}")
    return path


def to_ai_digest_file(summary: Dict) -> str:
    """把 AI 摘要落成 .md 文件（便于直接投喂）"""
    path = export_path("知识图谱AI摘要", "md")
    with open(path, "w", encoding="utf-8") as f:
        f.write(to_ai_digest(summary))
    return path
